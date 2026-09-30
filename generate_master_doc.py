import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_image_with_caption(doc, img_path, caption_text, width_inches=5.6):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_img.add_run()
        run.add_picture(img_path, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run(f"Gambar: {caption_text}")
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(100, 100, 100)
        doc.add_paragraph()

doc = docx.Document()
sec = doc.sections[0]
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)

# ==========================================
# COVER / HEADER
# ==========================================
title_p = doc.add_paragraph()
title_r = title_p.add_run('DOKUMENTASI LENGKAP & SPESIFIKASI SISTEM (PRD & ARSITEKTUR)')
title_r.bold = True
title_r.font.size = Pt(20)
title_r.font.color.rgb = RGBColor(31, 78, 120)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub_p = doc.add_paragraph()
sub_r = sub_p.add_run('Nexus Read / LibLog: Web Perpustakaan Cerdas, Gamifikasi Streak, Diary Bacaan (Letterboxd-Style) & Machine Learning')
sub_r.italic = True
sub_r.font.size = Pt(11)
sub_r.font.color.rgb = RGBColor(89, 89, 89)
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()

# ==========================================
# BAB 1: INFORMASI DOKUMEN & RINGKASAN EKSEKUTIF
# ==========================================
h1 = doc.add_heading('BAB 1: Ringkasan Eksekutif & Target Pengguna', level=1)
h1.runs[0].font.color.rgb = RGBColor(31, 78, 120)

table = doc.add_table(rows=5, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
data_info = [
    ('Nama Produk', 'Nexus Read / LibLog (Smart Web Library & Reading Habit Tracker)'),
    ('Tipe Produk', 'Web Application (Sirkulasi Perpustakaan + Reading Diary + ML Estimator)'),
    ('Versi Dokumen', 'v1.0 (Master Comprehensive Document)'),
    ('Target Rilis', 'Fase 1 / MVP (Minimum Viable Product)'),
    ('Status Proyek', 'Disetujui untuk Tahap Implementasi (Approved for Development)')
]
for row_idx, (k, v) in enumerate(data_info):
    row = table.rows[row_idx]
    row.cells[0].width = Inches(1.8)
    row.cells[0].text = k
    row.cells[0].paragraphs[0].runs[0].bold = True
    set_cell_background(row.cells[0], 'F2F4F7')
    set_cell_margins(row.cells[0], 100, 100, 150, 150)
    
    row.cells[1].width = Inches(4.7)
    row.cells[1].text = v
    set_cell_margins(row.cells[1], 100, 100, 150, 150)

doc.add_paragraph()

p_vis = doc.add_paragraph()
p_vis.add_run('1.1 Visi Produk:\n').bold = True
p_vis.add_run('Mengubah sistem perpustakaan konvensional menjadi platform web interaktif modern yang tidak hanya memfasilitasi sirkulasi buku (pinjam-kembali), tetapi juga memberdayakan pengguna untuk membangun kebiasaan membaca harian melalui sistem Gamifikasi Streak, pencatatan Reading Diary (ala Letterboxd), dan estimasi durasi penyelesaian buku personal berbasis Machine Learning.')

doc.add_paragraph().add_run('1.2 Target Pengguna & Persona:').bold = True
personas = [
    ('Mahasiswa / Pembaca Aktif (The Avid Reader)', 'Meminjam buku mandiri, melacak progress buku yang sedang/sudah dibaca, memberikan rating/review ala Letterboxd, serta menjaga streak harian.'),
    ('Pembaca Kasual / Pemula (The Habit Builder)', 'Membutuhkan estimasi realistis berapa hari buku selesai dibaca agar terhindar dari denda dan termotivasi oleh visual streak harian.'),
    ('Pustakawan / Administrator', 'Mengelola ketersediaan stok buku fisik, memantau sirkulasi peminjaman, serta melacak status pengembalian otomatis.')
]
for title, desc in personas:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(f'{title}: ').bold = True
    p.add_run(desc)

# ==========================================
# BAB 2: KEBUTUHAN FUNGSIONAL (PRD)
# ==========================================
h2 = doc.add_heading('BAB 2: Kebutuhan Fungsional Sistem (Functional Requirements)', level=1)
h2.runs[0].font.color.rgb = RGBColor(31, 78, 120)

mods = [
    ('2.1 Modul 1: Autentikasi & Manajemen Akun', [
        ('FR-1.1 Register & Login', 'Registrasi dan login dengan identitas NIM/Email & password terenkripsi aman.'),
        ('FR-1.2 Session Management', 'Sesi aman berbasis token JWT.'),
        ('FR-1.3 Logout', 'Fitur keluar akun aman dengan menghapus token sesi pengguna.'),
        ('FR-1.4 Dashboard Profil', 'Menampilkan total buku dibaca, streak aktif, ulasan, dan lencana badge.')
    ]),
    ('2.2 Modul 2: Sirkulasi Perpustakaan (Core Library System)', [
        ('FR-2.1 Katalog & Pencarian Buku', 'Pencarian buku berdasarkan judul, pengarang, genre, dan jumlah halaman beserta status ketersediaan stok fisik.'),
        ('FR-2.2 Pinjam Buku Baru', 'Meminjam buku yang tersedia dengan pencatatan tanggal pinjam dan tanggal jatuh tempo (due date).'),
        ('FR-2.3 Daftar Buku yang Sedang Dipinjam', 'Halaman khusus yang menampilkan buku aktif di tangan pengguna, sisa hari pinjaman, dan progress baca.'),
        ('FR-2.4 Kembalikan Buku', 'Pengembalian buku dengan deteksi otomatis denda keterlambatan dan pop-up pembuatan ulasan diary.')
    ]),
    ('2.3 Modul 3: Reading Diary & Review (Letterboxd for Books)', [
        ('FR-3.1 Reading Log Entry', 'Mencatat tanggal selesai, rating bintang (0.5 s.d. 5.0 bintang), ulasan teks, dan catatan kutipan favorit.'),
        ('FR-3.2 Linimasa Riwayat Visual', 'Tampilan grid poster cover buku yang pernah dibaca sesuai linimasa bulan/tahun (Letterboxd style).'),
        ('FR-3.3 Rak Koleksi (Custom Lists)', 'Pengelompokan buku personal (misal: "Top 10 Buku Favorit", "Buku Wajib Baca 2026").')
    ]),
    ('2.4 Modul 4: Habit Tracker & Streak Gamification', [
        ('FR-4.1 Daily Check-In', 'Pengguna mencatat aktivitas progres halaman yang dibaca hari ini.'),
        ('FR-4.2 Dynamic Streak Counter', 'Sistem menghitung hari berturut-turut membaca aktif dengan indikator api interaktif.'),
        ('FR-4.3 Badges & Achievements', 'Lencana apresiasi atas konsistensi membaca (misal: "7-Day Streak", "Master Reader 500 Pages").')
    ]),
    ('2.5 Modul 5: Machine Learning Engine (Estimasi Durasi Membaca)', [
        ('FR-5.1 Prediksi Hari Selesai Membaca', 'Model regresi memprediksi durasi (jumlah hari) penyelesaian buku yang dipinjam.'),
        ('FR-5.2 Fitur yang Dipelajari', 'Jumlah halaman buku, genre/kategori, dan kecepatan historis baca pengguna (pages/day).'),
        ('FR-5.3 Peringatan Cerdas Jatuh Tempo', 'Rekomendasi cerdas jika estimasi baca pengguna melebihi durasi izin pinjam perpustakaan.')
    ])
]

for m_title, items in mods:
    sub_h = doc.add_heading(m_title, level=2)
    sub_h.runs[0].font.color.rgb = RGBColor(46, 117, 182)
    for code, desc in items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f'{code}: ').bold = True
        p.add_run(desc)

# ==========================================
# BAB 3: DIAGRAM ALUR PENGGUNA (USER FLOW)
# ==========================================
h3 = doc.add_heading('BAB 3: Alur Pengguna (User Flow Diagram)', level=1)
h3.runs[0].font.color.rgb = RGBColor(31, 78, 120)

p_flow_intro = doc.add_paragraph()
p_flow_intro.add_run('Diagram di bawah ini menggambarkan perjalanan pengguna (user journey) mulai dari membuka website, mencari dan meminjam buku dengan estimasi ML, mencatat streak harian, hingga mengembalikan buku dan mengisi ulasan diary:')

add_image_with_caption(doc, 'dokumentasi/assets/user_flow.png', 'Diagram Alur Pengguna (User Flow Interaktif)', width_inches=4.8)

# ==========================================
# BAB 4: ARSITEKTUR SISTEM (SYSTEM ARCHITECTURE)
# ==========================================
h4 = doc.add_heading('BAB 4: Arsitektur Sistem (High-Level Architecture)', level=1)
h4.runs[0].font.color.rgb = RGBColor(31, 78, 120)

add_image_with_caption(doc, 'dokumentasi/assets/architecture.png', 'Arsitektur 3-Tier + Machine Learning Service Layer', width_inches=5.8)

p_arch_desc = doc.add_paragraph()
p_arch_desc.add_run('Rincian Lapisan Arsitektur:\n').bold = True
p_arch_desc.add_run('1. Presentation Layer (Frontend): Dibangun dengan Next.js / React + Tailwind CSS untuk antarmuka interaktif, poster grid cover buku ala Letterboxd, dan animasi streak harian.\n')
p_arch_desc.add_run('2. Application Layer (Backend API): Dibangun dengan FastAPI / Python yang mengelola autentikasi JWT, sirkulasi peminjaman, kalkulasi denda keterlambatan, dan logika streak.\n')
p_arch_desc.add_run('3. ML Service Layer: Model Scikit-Learn berbasis regresi yang dimuat langsung ke memori untuk menghasilkan estimasi waktu baca secara real-time (<10ms).\n')
p_arch_desc.add_run('4. Data Layer: Basis data relasional (PostgreSQL / SQLite) untuk menyimpan data user, katalog buku, transaksi pinjaman, dan log ulasan.')

# ==========================================
# BAB 5: FLOW PROSES BISNIS & SEQUENCE DIAGRAM
# ==========================================
h5 = doc.add_heading('BAB 5: Rincian Alur Proses Bisnis & Sequence Diagram', level=1)
h5.runs[0].font.color.rgb = RGBColor(31, 78, 120)

doc.add_heading('5.1 Alur Peminjaman Buku + Estimasi Cerdas ML', level=2).runs[0].font.color.rgb = RGBColor(46, 117, 182)
add_image_with_caption(doc, 'dokumentasi/assets/flow_pinjam.png', 'Flowchart Peminjaman Buku dengan Estimasi ML', width_inches=4.8)

doc.add_heading('5.2 Alur Pelacak Kebiasaan & Streak Harian', level=2).runs[0].font.color.rgb = RGBColor(46, 117, 182)
add_image_with_caption(doc, 'dokumentasi/assets/flow_streak.png', 'Flowchart Daily Check-In & Streak Engine', width_inches=4.8)

doc.add_heading('5.3 Alur Pengembalian Buku & Diary Bacaan (Letterboxd-style)', level=2).runs[0].font.color.rgb = RGBColor(46, 117, 182)
add_image_with_caption(doc, 'dokumentasi/assets/flow_kembali.png', 'Flowchart Pengembalian Buku & Modal Ulasan Diary', width_inches=4.8)

doc.add_heading('5.4 Diagram Sekuensial Interaksi Sistem (Sequence Diagram)', level=2).runs[0].font.color.rgb = RGBColor(46, 117, 182)
add_image_with_caption(doc, 'dokumentasi/assets/sequence_diagram.png', 'Sequence Diagram Interaksi Client - API - ML - Database', width_inches=5.8)

# ==========================================
# BAB 6: SKEMA BASIS DATA (ERD)
# ==========================================
h6 = doc.add_heading('BAB 6: Desain & Skema Basis Data (ERD)', level=1)
h6.runs[0].font.color.rgb = RGBColor(31, 78, 120)

add_image_with_caption(doc, 'dokumentasi/assets/erd_diagram.png', 'Entity-Relationship Diagram (ERD)', width_inches=5.8)

db_table = doc.add_table(rows=5, cols=3)
db_table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['Nama Tabel', 'Atribut Kunci', 'Keterangan']
for i, h in enumerate(headers):
    cell = db_table.rows[0].cells[i]
    cell.text = h
    cell.paragraphs[0].runs[0].bold = True
    set_cell_background(cell, '1F4E78')
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    set_cell_margins(cell, 120, 120, 120, 120)

db_rows = [
    ('Users', 'user_id, nim_or_email, name, password_hash, current_streak, max_streak, last_checkin', 'Menyimpan akun pengguna, keamanan auth, dan status streak'),
    ('Books', 'book_id, title, author, category, total_pages, stock_qty, cover_url', 'Menyimpan master katalog buku perpustakaan'),
    ('Borrow_Transactions', 'tx_id, user_id, book_id, borrow_date, due_date, return_date, status, ml_predicted_days', 'Mencatat sirkulasi peminjaman aktif dan jatuh tempo'),
    ('Reading_Logs', 'log_id, user_id, book_id, rating, review_text, finished_date, days_spent', 'Mencatat riwayat buku dibaca, rating & ulasan (Letterboxd-style)')
]

for row_idx, (t_name, attrs, desc) in enumerate(db_rows, start=1):
    row = db_table.rows[row_idx]
    row.cells[0].text = t_name
    row.cells[0].paragraphs[0].runs[0].bold = True
    row.cells[1].text = attrs
    row.cells[2].text = desc
    for c in row.cells:
        set_cell_margins(c, 100, 100, 120, 120)

doc.add_paragraph()

# ==========================================
# BAB 7: PIPELINE MACHINE LEARNING
# ==========================================
h7 = doc.add_heading('BAB 7: Desain & Pipeline Machine Learning', level=1)
h7.runs[0].font.color.rgb = RGBColor(31, 78, 120)

add_image_with_caption(doc, 'dokumentasi/assets/ml_pipeline.png', 'Pipeline Machine Learning (Reading Pace Estimator)', width_inches=5.8)

p_ml_desc = doc.add_paragraph()
p_ml_desc.add_run('Tahapan Model Estimasi Waktu Baca:\n').bold = True
p_ml_desc.add_run('1. Data Harvesting: Mengumpulkan durasi penyelesaian buku (days_spent) dan rata-rata halaman per hari dari transaksi terdahulu.\n')
p_ml_desc.add_run('2. Feature Engineering: Ekstraksi fitur [total_pages, category_encoded, user_historical_speed, day_of_week].\n')
p_ml_desc.add_run('3. Model Training: Pelatihan model regresi (Random Forest / Ridge Regression) untuk menghasilkan prediksi hari baca.\n')
p_ml_desc.add_run('4. Cold-Start Solution: Bagi pengguna baru, sistem menggunakan default kecepatan baca global (25 hal/hari untuk fiksi, 15 hal/hari untuk non-fiksi).')

# ==========================================
# BAB 8: ROADMAP IMPLEMENTASI
# ==========================================
h8 = doc.add_heading('BAB 8: Roadmap Implementasi & Milestone', level=1)
h8.runs[0].font.color.rgb = RGBColor(31, 78, 120)

milestones = [
    ('Fase 1: Core & Auth MVP', 'Autentikasi (Register, Login, Logout), Katalog Buku, Fitur Pinjam Buku Baru, List Buku Sedang Dipinjam, dan Kembalikan Buku.'),
    ('Fase 2: Reading Diary & Gamifikasi', 'Reading Diary ala Letterboxd (Rating, Review, Grid Cover View) dan Sistem Daily Check-In + Streak Counter.'),
    ('Fase 3: Machine Learning Engine', 'Pengumpulan dataset log bacaan, integrasi model regresi estimasi waktu baca, serta rekomendasi cerdas pada antarmuka web.'),
    ('Fase 4: Testing & Deployment', 'Pengujian menyeluruh, optimalisasi UI/UX responsif, dan peluncuran web aplikasi.')
]

for phase, desc in milestones:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(f'{phase}: ').bold = True
    p.add_run(desc)

# Save Master Docx
out_path = 'dokumentasi/DOKUMENTASI_MASTER_WEB_PERPUSTAKAAN_NEXUSREAD.docx'
try:
    doc.save(out_path)
    print(f"Master Document successfully created: {out_path}")
except Exception as e:
    print(f"Error saving: {e}")
